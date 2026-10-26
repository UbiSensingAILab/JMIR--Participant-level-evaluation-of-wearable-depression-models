#!/usr/bin/env python3
# REB25-1587_MOD1 — research data module; no PII in logs
"""Shared confusion-matrix and case-by-case export helper for the Ref 2
Ablation Study tab's tables 6, 7, 8 and 8b.

Locked positive-class convention (see <positive_class_convention> in the
quick task 260903-dpq plan, consistent with every ROC-AUC/accuracy figure
already on this tab): `label == 1` means the participant-segment's PHQ-9
score is >= 10. Predicted label is `P(label==1) >= DECISION_THRESHOLD`
(0.5, the same boundary `pipeline.predict` and `oof_proba >= 0.5` already
use everywhere on this tab). TP = actual 1 & predicted 1; FP = actual 0 &
predicted 1; TN = actual 0 & predicted 0; FN = actual 1 & predicted 0.

This module NEVER fits, refits, or re-thresholds anything — it only counts
and formats predictions handed to it by the caller. Every number it
produces is derived from arrays a sibling script already computed
(`run_single_split_full`'s held-out arrays, `run_cv`'s pooled
out-of-fold arrays); no new model is trained here.

Quick task 260908-l0q adds a second, purely additive convention below the
original one. This module now serves TWO conventions side by side: the
original locked binary one just described (label 0/1, threshold-derived
predictions via `DECISION_THRESHOLD`, TP/FP/TN/FN cell names) used by the
four HRV/Depresjon call sites above, and a new additive multi-class one
(`multiclass_case_rows` / `multiclass_cases_markdown`) used by the
RESILIENT participant-level classifier, which has no positive class, no
probability threshold, and therefore no TP/FP/TN/FN cell-name concept —
just a plain actual-class-vs-predicted-class pair per case.
"""

from __future__ import annotations

import math
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Font

# Fixed, not tunable — the threshold the already-published accuracy figures
# on this tab were computed at (`pooled_metrics` uses `oof_proba >= 0.5`;
# `run_single_split_full` uses `pipeline.predict`, the same boundary for a
# binary classifier). Not a free parameter of this module.
POSITIVE_LABEL = 1
DECISION_THRESHOLD = 0.5

_ACCURACY_TOLERANCE = 5e-5


def cell_label(actual: int, predicted: int) -> str:
    """Returns one of the four two-letter confusion-cell names."""
    if actual == 1 and predicted == 1:
        return "TP"
    if actual == 0 and predicted == 1:
        return "FP"
    if actual == 0 and predicted == 0:
        return "TN"
    if actual == 1 and predicted == 0:
        return "FN"
    raise ValueError(f"actual/predicted must be 0/1, got actual={actual!r} predicted={predicted!r}")


def confusion_counts(
    y_true: np.ndarray, proba: np.ndarray, expected_accuracy: float | None = None
) -> dict:
    """Counts TP/FP/TN/FN (and derived rates) from actual labels and
    predicted-positive probabilities. Predictions are derived as
    `proba >= DECISION_THRESHOLD` — never a free-standing prediction array,
    so the threshold this module uses is always explicit and traceable.

    When `expected_accuracy` is supplied, asserts the counted accuracy
    ((tp+tn)/n) agrees with it to within 5e-5 — this is the mechanism that
    ties these new counts to each script's already-published accuracy
    figure. Guards a zero denominator in sensitivity/specificity/precision
    by returning NaN rather than raising, mirroring pooled_metrics's
    NaN-for-degenerate-case convention in train_hrv_phq9_baseline.py.
    """
    y_true = np.asarray(y_true).astype(int)
    proba = np.asarray(proba).astype(float)
    assert len(y_true) == len(proba), (
        f"y_true and proba must be the same length: {len(y_true)} vs. {len(proba)}"
    )

    predicted = (proba >= DECISION_THRESHOLD).astype(int)
    n = int(len(y_true))

    tp = int(((y_true == 1) & (predicted == 1)).sum())
    fp = int(((y_true == 0) & (predicted == 1)).sum())
    tn = int(((y_true == 0) & (predicted == 0)).sum())
    fn = int(((y_true == 1) & (predicted == 0)).sum())

    assert tp + fp + tn + fn == n, (
        f"confusion cells must sum to n: tp={tp} fp={fp} tn={tn} fn={fn} "
        f"sum={tp + fp + tn + fn} n={n}"
    )

    n_actual_positive = tp + fn
    n_actual_negative = tn + fp

    accuracy = (tp + tn) / n if n else float("nan")
    sensitivity = tp / n_actual_positive if n_actual_positive > 0 else float("nan")
    specificity = tn / n_actual_negative if n_actual_negative > 0 else float("nan")
    precision = tp / (tp + fp) if (tp + fp) > 0 else float("nan")

    if expected_accuracy is not None:
        assert abs(accuracy - expected_accuracy) <= _ACCURACY_TOLERANCE, (
            f"confusion_counts accuracy {accuracy:.6f} disagrees with the "
            f"already-published accuracy {expected_accuracy:.6f} by more than "
            f"{_ACCURACY_TOLERANCE} — HALT and report the real numbers, do not "
            "adjust either value to match the other."
        )

    return {
        "tp": tp,
        "fp": fp,
        "tn": tn,
        "fn": fn,
        "n": n,
        "n_actual_positive": n_actual_positive,
        "n_actual_negative": n_actual_negative,
        "sensitivity": sensitivity,
        "specificity": specificity,
        "precision": precision,
        "accuracy": accuracy,
    }


def case_rows(
    device_ids: np.ndarray,
    y_true: np.ndarray,
    proba_by_model: dict[str, np.ndarray],
    folds: np.ndarray | None = None,
) -> list[dict]:
    """One row per test window: index (1-based), device_id, fold (int or
    None), actual, and per model in proba_by_model's order: '<name> pred',
    '<name> P(pos)', '<name> cell'."""
    device_ids = np.asarray(device_ids)
    y_true = np.asarray(y_true).astype(int)
    n = len(y_true)
    assert len(device_ids) == n, f"device_ids length {len(device_ids)} != y_true length {n}"
    for name, proba in proba_by_model.items():
        assert len(proba) == n, (
            f"proba_by_model[{name!r}] length {len(proba)} != y_true length {n}"
        )
    if folds is not None:
        folds = np.asarray(folds)
        assert len(folds) == n, f"folds length {len(folds)} != y_true length {n}"

    rows: list[dict] = []
    for i in range(n):
        row: dict = {
            "index": i + 1,
            "device_id": device_ids[i],
            "fold": int(folds[i]) if folds is not None else None,
            "actual": int(y_true[i]),
        }
        for name, proba in proba_by_model.items():
            p = float(proba[i])
            pred = int(p >= DECISION_THRESHOLD)
            row[f"{name} pred"] = pred
            row[f"{name} P(pos)"] = p
            row[f"{name} cell"] = cell_label(int(y_true[i]), pred)
        rows.append(row)
    return rows


def confusion_markdown(counts_by_model: dict[str, dict]) -> list[str]:
    """Emits a markdown table: one row per model, header exactly as
    specified in the plan."""
    lines = [
        "| Model | TP | FP | TN | FN | N | Sensitivity | Specificity | Precision | Accuracy (TP+TN)/N |",
        "|---|---|---|---|---|---|---|---|---|---|",
    ]

    def _fmt_rate(v: float) -> str:
        return "NaN" if isinstance(v, float) and math.isnan(v) else f"{v:.4f}"

    for name, counts in counts_by_model.items():
        lines.append(
            f"| {name} | {counts['tp']} | {counts['fp']} | {counts['tn']} | {counts['fn']} | "
            f"{counts['n']} | {_fmt_rate(counts['sensitivity'])} | "
            f"{_fmt_rate(counts['specificity'])} | {_fmt_rate(counts['precision'])} | "
            f"{_fmt_rate(counts['accuracy'])} |"
        )
    return lines


def cases_markdown(rows: list[dict], model_names: list[str], include_fold: bool) -> list[str]:
    """Emits a markdown table with columns #, deviceId, [Fold], Actual, then
    per model in `model_names` order: pred, P(pos), cell."""
    header_cells = ["#", "deviceId"]
    if include_fold:
        header_cells.append("Fold")
    header_cells.append("Actual")
    for name in model_names:
        header_cells += [f"{name} pred", f"{name} P(pos)", f"{name} cell"]

    lines = [
        "| " + " | ".join(header_cells) + " |",
        "|" + "---|" * len(header_cells),
    ]

    for row in rows:
        cells = [str(row["index"]), str(row["device_id"])]
        if include_fold:
            cells.append(str(row["fold"]))
        cells.append(str(row["actual"]))
        for name in model_names:
            cells.append(str(row[f"{name} pred"]))
            cells.append(f"{row[f'{name} P(pos)']:.4f}")
            cells.append(row[f"{name} cell"])
        lines.append("| " + " | ".join(cells) + " |")

    return lines


def write_cases_workbook(path: str | Path, sheets: dict[str, tuple[list[str], list[list]]]) -> dict[str, int]:
    """Writes an .xlsx workbook. `sheets` is an ordered mapping of sheet
    name -> (header_list, row_lists). Returns per-sheet data-row counts."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    row_counts: dict[str, int] = {}
    for sheet_name, (header, data_rows) in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for data_row in data_rows:
            ws.append(data_row)
        ws.freeze_panes = "A2"
        for col_idx, col_name in enumerate(header, start=1):
            width = max(12, min(40, len(str(col_name)) + 4))
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = width
        row_counts[sheet_name] = len(data_rows)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return row_counts


# ---------------------------------------------------------------------------
# Additive multi-class convention (quick task 260908-l0q) — no positive
# class, no probability threshold, no TP/FP/TN/FN cell names. Used by the
# RESILIENT participant-level PHQ-9/GDS-15 classifier. Nothing above this
# point is modified by these additions.
# ---------------------------------------------------------------------------

def multiclass_case_rows(
    case_ids: np.ndarray,
    y_true: np.ndarray,
    pred_by_model: dict[str, np.ndarray],
    folds: np.ndarray | None = None,
    class_labels: tuple | list | None = None,
) -> list[dict]:
    """One row per case: index (1-based), case_id, fold (int or None),
    actual (int), actual_label, and per model in pred_by_model's order:
    '<name> pred' (int), '<name> pred_label'. `class_labels` is an optional
    sequence indexed by class integer; when it is None, the label fields
    repeat the integer as a string. Every value in y_true and in each
    prediction array must be a valid index into class_labels when it is
    supplied. Emits no probability column and no cell name — this
    convention has no positive class, so a plain actual/predicted pair is
    the whole payload."""
    case_ids = np.asarray(case_ids)
    y_true = np.asarray(y_true).astype(int)
    n = len(y_true)
    assert len(case_ids) == n, f"case_ids length {len(case_ids)} != y_true length {n}"
    for name, preds in pred_by_model.items():
        assert len(preds) == n, (
            f"pred_by_model[{name!r}] length {len(preds)} != y_true length {n}"
        )
    if folds is not None:
        folds = np.asarray(folds)
        assert len(folds) == n, f"folds length {len(folds)} != y_true length {n}"

    def _label(v: int) -> str:
        if class_labels is not None:
            assert 0 <= v < len(class_labels), (
                f"class index {v} out of range for class_labels of length {len(class_labels)}"
            )
            return str(class_labels[v])
        return str(v)

    rows: list[dict] = []
    for i in range(n):
        row: dict = {
            "index": i + 1,
            "case_id": case_ids[i],
            "fold": int(folds[i]) if folds is not None else None,
            "actual": int(y_true[i]),
            "actual_label": _label(int(y_true[i])),
        }
        for name, preds in pred_by_model.items():
            p = int(preds[i])
            row[f"{name} pred"] = p
            row[f"{name} pred_label"] = _label(p)
        rows.append(row)
    return rows


def multiclass_cases_markdown(
    rows: list[dict],
    model_names: list[str],
    include_fold: bool,
    id_header: str = "caseId",
    prefix_model_names: bool = True,
) -> list[str]:
    """Emits a markdown table with columns #, id_header, [Fold], Actual,
    then per model either '<name> predicted' (prefix_model_names=True) or
    plain 'Predicted' (prefix_model_names=False — requires exactly one
    model, since unprefixed columns are ambiguous with more than one
    model). Each class cell renders as the integer followed by its label
    in parentheses when the label differs from the bare integer string,
    otherwise the bare integer."""
    if not prefix_model_names:
        assert len(model_names) == 1, (
            "prefix_model_names=False requires exactly one model — unprefixed "
            "columns are ambiguous with more than one model"
        )

    header_cells = ["#", id_header]
    if include_fold:
        header_cells.append("Fold")
    header_cells.append("Actual")
    if prefix_model_names:
        for name in model_names:
            header_cells.append(f"{name} predicted")
    else:
        header_cells.append("Predicted")

    lines = [
        "| " + " | ".join(header_cells) + " |",
        "|" + "---|" * len(header_cells),
    ]

    def _cell(value: int, label: str) -> str:
        s = str(value)
        return f"{s} ({label})" if label != s else s

    for row in rows:
        cells = [str(row["index"]), str(row["case_id"])]
        if include_fold:
            cells.append(str(row["fold"]))
        cells.append(_cell(row["actual"], row["actual_label"]))
        if prefix_model_names:
            for name in model_names:
                cells.append(_cell(row[f"{name} pred"], row[f"{name} pred_label"]))
        else:
            name = model_names[0]
            cells.append(_cell(row[f"{name} pred"], row[f"{name} pred_label"]))
        lines.append("| " + " | ".join(cells) + " |")

    return lines
